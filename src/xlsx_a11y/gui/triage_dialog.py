"""Visual interactive remediation dialog for Excel accessibility barriers."""
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import openpyxl
from PySide6.QtWidgets import (
    QCheckBox,
    QDialog,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QPushButton,
    QVBoxLayout,
)


class TriageDialog(QDialog):
    def __init__(
        self,
        xlsx_path: Union[Path, str],
        findings: List[Any],
        parent=None,
    ):
        super().__init__(parent)
        self.xlsx_path = Path(xlsx_path)
        self.wb = openpyxl.load_workbook(self.xlsx_path, data_only=False)

        # Filter triageable findings
        self.findings = []
        for f in findings:
            rule_id = f.rule_id if hasattr(f, "rule_id") else f.get("rule_id", "")
            if rule_id in (
                "image-alt-missing",
                "chart-alt-missing",
                "title-missing",
                "sheet-name-default",
            ):
                self.findings.append(f)

        self.current_idx = 0
        self.items_modified = 0

        self.setWindowTitle("Interactive Remediation Triage")
        self.resize(600, 320)
        self._init_ui()
        self._load_current()

    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(12)

        self.lbl_progress = QLabel("Barrier 1 of 1")
        self.lbl_progress.setStyleSheet("font-weight: 600; color: #16a34a;")
        layout.addWidget(self.lbl_progress)

        self.lbl_location = QLabel("Location:")
        self.lbl_location.setStyleSheet("font-weight: 600; color: #334155;")
        layout.addWidget(self.lbl_location)

        self.lbl_desc = QLabel("Description:")
        self.lbl_desc.setWordWrap(True)
        layout.addWidget(self.lbl_desc)

        # Input row
        self.lbl_input_prompt = QLabel("Descriptive Alternative Text / Value:")
        layout.addWidget(self.lbl_input_prompt)

        self.txt_input = QLineEdit()
        layout.addWidget(self.txt_input)

        self.chk_decorative = QCheckBox("Mark as Decorative (Purely aesthetic, ignore in screen readers)")
        self.chk_decorative.toggled.connect(self._on_decorative_toggled)
        layout.addWidget(self.chk_decorative)

        layout.addStretch()

        # Action Buttons
        btn_layout = QHBoxLayout()
        self.btn_skip = QPushButton("Skip")
        self.btn_skip.clicked.connect(self.skip_current)

        self.btn_apply = QPushButton("Apply & Next")
        self.btn_apply.setStyleSheet("background-color: #16a34a; color: white; font-weight: 600;")
        self.btn_apply.clicked.connect(self.apply_current)

        btn_layout.addStretch()
        btn_layout.addWidget(self.btn_skip)
        btn_layout.addWidget(self.btn_apply)
        layout.addLayout(btn_layout)

    def _on_decorative_toggled(self, checked: bool):
        self.txt_input.setEnabled(not checked)
        if checked:
            self.txt_input.clear()

    def _load_current(self):
        if not self.findings or self.current_idx >= len(self.findings):
            self._finish()
            return

        f = self.findings[self.current_idx]
        tot = len(self.findings)
        rule_id = f.rule_id if hasattr(f, "rule_id") else f.get("rule_id", "")
        location = f.location if hasattr(f, "location") else f.get("location", "")
        desc = f.description if hasattr(f, "description") else f.get("description", "")

        self.lbl_progress.setText(f"Barrier {self.current_idx + 1} of {tot}")
        self.lbl_location.setText(f"Location: {location}")
        self.lbl_desc.setText(desc)

        self.chk_decorative.setChecked(False)
        self.txt_input.clear()
        self.txt_input.setEnabled(True)

        if rule_id in ("image-alt-missing", "chart-alt-missing"):
            self.lbl_input_prompt.setText("Descriptive Alternative Text / Title:")
            self.chk_decorative.setVisible(rule_id == "image-alt-missing")
        elif rule_id == "title-missing":
            self.lbl_input_prompt.setText("Document Title (WCAG 2.4.2):")
            self.chk_decorative.setVisible(False)
        elif rule_id == "sheet-name-default":
            self.lbl_input_prompt.setText("Descriptive Sheet Tab Name (WCAG 2.4.6):")
            self.chk_decorative.setVisible(False)

    def skip_current(self):
        self.current_idx += 1
        self._load_current()

    def apply_current(self):
        if not self.findings or self.current_idx >= len(self.findings):
            return

        f = self.findings[self.current_idx]
        rule_id = f.rule_id if hasattr(f, "rule_id") else f.get("rule_id", "")
        location = f.location if hasattr(f, "location") else f.get("location", "")
        val = self.txt_input.text().strip()
        is_dec = self.chk_decorative.isChecked()

        if rule_id == "title-missing" and val:
            self.wb.properties.title = val
            self.items_modified += 1

        elif rule_id == "sheet-name-default" and val:
            if location in self.wb.sheetnames:
                self.wb[location].title = val[:31]
                self.items_modified += 1

        elif rule_id == "chart-alt-missing" and val:
            if "!" in location:
                sheet_name = location.split("!", 1)[0]
                if sheet_name in self.wb.sheetnames:
                    ws = self.wb[sheet_name]
                    if hasattr(ws, "_charts") and ws._charts:
                        ws._charts[0].title = val
                        self.items_modified += 1

        elif rule_id == "image-alt-missing":
            if "!" in location:
                sheet_name = location.split("!", 1)[0]
                if sheet_name in self.wb.sheetnames:
                    ws = self.wb[sheet_name]
                    if hasattr(ws, "_images") and ws._images:
                        img = ws._images[0]
                        img.descr = "" if is_dec else val
                        self.items_modified += 1

        self.current_idx += 1
        self._load_current()

    def _finish(self):
        if self.items_modified > 0:
            out_p = self.xlsx_path.parent / f"{self.xlsx_path.stem}-triaged.xlsx"
            self.wb.save(out_p)
        self.accept()
