"""Interactive terminal triage workflow for Excel workbook remediation."""
from pathlib import Path
from typing import Any, Callable, Optional, Union
import openpyxl

from xslx_a11y.audit import audit_file
from xslx_a11y.immutability import calculate_sha256, verify_immutability


def run_interactive_triage(
    in_path: Union[str, Path],
    out_path: Optional[Union[str, Path]] = None,
    input_func: Optional[Callable[[str], str]] = None,
    print_func: Optional[Callable[..., None]] = None,
) -> int:
    """Interactively walks the user through author-intent accessibility barriers."""
    if input_func is None:
        input_func = input
    if print_func is None:
        print_func = print

    in_p = Path(in_path).resolve()
    if not in_p.exists() or not in_p.is_file():
        raise FileNotFoundError(f"Input file not found: {in_p}")

    out_p = Path(out_path).resolve() if out_path else in_p.parent / f"{in_p.stem}-triaged.xlsx"
    if in_p == out_p:
        raise ValueError("Source and triage output file must be different paths.")

    sha_before = calculate_sha256(in_p)
    wb = openpyxl.load_workbook(in_p, data_only=False)
    audit_res = audit_file(in_p)
    findings = audit_res.get("findings", [])

    items_triaged = 0
    print_func(f"\n=== xslx-a11y Interactive Accessibility Triage: {in_p.name} ===\n")

    for f in findings:
        rule_id = f.rule_id if hasattr(f, "rule_id") else f.get("rule_id", "")
        location = f.location if hasattr(f, "location") else f.get("location", "")
        desc = f.description if hasattr(f, "description") else f.get("description", "")

        if rule_id == "title-missing":
            print_func("\n[Document Missing Title]")
            print_func(f"Location: {location}")
            print_func(f"Issue:    {desc}")
            ans = input_func("Enter a title for this workbook (or 's' to skip): ").strip()
            if ans.lower() != "s" and ans:
                wb.properties.title = ans
                items_triaged += 1
                print_func(f"-> Set document title: '{ans}'")

        elif rule_id == "sheet-name-default":
            print_func("\n[Default Sheet Tab Name]")
            print_func(f"Sheet:    {location}")
            print_func(f"Issue:    {desc}")
            ans = input_func("Enter a new descriptive name for this sheet (or 's' to skip): ").strip()
            if ans.lower() != "s" and ans:
                if location in wb.sheetnames:
                    ws = wb[location]
                    ws.title = ans[:31]
                    items_triaged += 1
                    print_func(f"-> Renamed sheet '{location}' to '{ans[:31]}'")

        elif rule_id == "chart-alt-missing":
            print_func("\n[Chart Missing Title / Alt Text]")
            print_func(f"Location: {location}")
            print_func(f"Issue:    {desc}")
            ans = input_func("Enter descriptive title for chart (or 's' to skip): ").strip()
            if ans.lower() != "s" and ans:
                # Location format: "{sheet}!chart[{idx}]"
                if "!" in location:
                    sheet_name, target = location.split("!", 1)
                    if sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                        if hasattr(ws, "_charts") and ws._charts:
                            ws._charts[0].title = ans
                            items_triaged += 1
                            print_func(f"-> Set chart title to: '{ans}'")

        elif rule_id == "image-alt-missing":
            print_func("\n[Image Missing Alt Text]")
            print_func(f"Location: {location}")
            print_func(f"Issue:    {desc}")
            ans = input_func("Enter alt text description (or 's' to skip): ").strip()
            if ans.lower() != "s" and ans:
                if "!" in location:
                    sheet_name, target = location.split("!", 1)
                    if sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                        if hasattr(ws, "_images") and ws._images:
                            ws._images[0].descr = ans
                            items_triaged += 1
                            print_func(f"-> Set image alt text: '{ans}'")

    out_p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_p)
    verify_immutability(in_p, sha_before)

    print_func(f"\nTriage complete! {items_triaged} items updated. Saved to: {out_p}\n")
    return items_triaged
