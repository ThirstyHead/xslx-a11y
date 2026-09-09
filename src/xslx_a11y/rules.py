"""Accessibility audit rules for Microsoft Excel (.xlsx) workbooks.

Evaluates workbooks against WCAG 2.1 AA criteria:
- 2.4.2 Page Titled (Workbook title)
- 2.4.6 Headings & Labels (Descriptive worksheet tab names)
- 1.3.1 Info & Relationships (Empty sheets, merged cells, table headers)
- 1.1.1 Non-text Content (Image and chart alternative text)
- 1.4.3 Contrast (Minimum) (Cell font vs fill contrast ratio)
- 2.4.4 Link Purpose (Descriptive hyperlink text)
"""
import re
from typing import List, Optional, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
import wcag_contrast_ratio

from xslx_a11y.findings import Finding

VAGUE_LINK_TEXTS = {
    "click here",
    "click here.",
    "here",
    "read more",
    "more",
    "link",
    "learn more",
    "details",
    "go",
    "view",
}


def hex_to_rgb_tuple(hex_str: str) -> Optional[Tuple[float, float, float]]:
    """Converts hex color string (6 or 8 chars ARGB/RGB) to normalized (r, g, b)."""
    if not hex_str or not isinstance(hex_str, str):
        return None
    s = hex_str.strip().lstrip("#")
    if len(s) == 8:
        # ARGB format from openpyxl - drop alpha
        s = s[2:]
    if len(s) != 6:
        return None
    try:
        r = int(s[0:2], 16) / 255.0
        g = int(s[2:4], 16) / 255.0
        b = int(s[4:6], 16) / 255.0
        return (r, g, b)
    except ValueError:
        return None


def check_workbook_title(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 2.4.2: Checks that the workbook has a meaningful title property."""
    findings = []
    title = getattr(wb.properties, "title", None)
    if not title or not str(title).strip():
        findings.append(
            Finding(
                rule_id="title-missing",
                sc="2.4.2",
                severity="critical",
                location="core.xml",
                description="Workbook title is missing or blank in document properties",
                evidence=f"title={repr(title)}",
                fixable=True,
                fix="Set sanitized document title in core properties",
            )
        )
    return findings


def check_sheet_names(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 2.4.6: Checks for default, non-descriptive worksheet names (e.g. Sheet1)."""
    findings = []
    pattern = re.compile(r"^Sheet\d+$", re.IGNORECASE)
    for ws in wb.worksheets:
        if pattern.match(ws.title.strip()):
            findings.append(
                Finding(
                    rule_id="sheet-name-default",
                    sc="2.4.6",
                    severity="moderate",
                    location=ws.title,
                    description=f"Worksheet '{ws.title}' uses a default, non-descriptive tab name",
                    evidence=f"title='{ws.title}'",
                    fixable=False,
                    fix="Rename worksheet tab with a concise, descriptive label",
                )
            )
    return findings


def check_empty_sheets(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.3.1: Flags empty worksheets that convey no information."""
    findings = []
    for ws in wb.worksheets:
        is_empty = True
        if ws.max_row > 1 or ws.max_column > 1:
            is_empty = False
        else:
            cell_val = ws.cell(row=1, column=1).value
            if cell_val is not None and str(cell_val).strip() != "":
                is_empty = False

        if is_empty:
            findings.append(
                Finding(
                    rule_id="sheet-empty",
                    sc="1.3.1",
                    severity="moderate",
                    location=ws.title,
                    description=f"Worksheet '{ws.title}' is completely empty",
                    evidence="max_row=1, max_column=1, cell(1,1)=None",
                    fixable=False,
                    fix="Delete unused empty sheet or add relevant content",
                )
            )
    return findings


def check_merged_cells(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.3.1 / 1.3.2: Flags merged cells that break screen reader grid navigation."""
    findings = []
    for ws in wb.worksheets:
        for rng in list(ws.merged_cells.ranges):
            findings.append(
                Finding(
                    rule_id="merged-cell",
                    sc="1.3.1",
                    severity="serious",
                    location=f"{ws.title}!{rng.coord}",
                    description=(
                        f"Merged cell range '{rng.coord}' disrupts screen reader row/column "
                        "association and linear navigation"
                    ),
                    evidence=f"merged_range='{rng.coord}'",
                    fixable=True,
                    fix="Unmerge cells and replicate content across previously merged cells",
                )
            )
    return findings


def _is_coord_in_table(min_r: int, min_c: int, max_r: int, max_c: int, table) -> bool:
    """Helper to check if bounding box is within an existing table."""
    try:
        from openpyxl.utils.cell import range_boundaries
        t_min_c, t_min_r, t_max_c, t_max_r = range_boundaries(table.ref)
        if None in (t_min_c, t_min_r, t_max_c, t_max_r):
            return False
        return (
            min_r >= t_min_r  # type: ignore[operator]
            and max_r <= t_max_r  # type: ignore[operator]
            and min_c >= t_min_c  # type: ignore[operator]
            and max_c <= t_max_c  # type: ignore[operator]
        )
    except Exception:
        return False


def check_table_headers(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.3.1: Verifies tables have explicit header rows and data ranges are structured."""
    findings = []
    for ws in wb.worksheets:
        tables = list(ws.tables.values())

        # Check existing tables for header row
        for table in tables:
            if table.headerRowCount is not None and table.headerRowCount == 0:
                loc = f"{ws.title}!{table.ref or table.displayName}"
                findings.append(
                    Finding(
                        rule_id="table-header-missing",
                        sc="1.3.1",
                        severity="serious",
                        location=loc,
                        description=f"Table '{table.displayName}' does not have a designated header row",
                        evidence=f"headerRowCount={table.headerRowCount}",
                        fixable=True,
                        fix="Enable header row on table",
                    )
                )

        # Check for contiguous data block without an official Excel Table
        non_empty_rows = []
        non_empty_cols = []
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is not None and str(cell.value).strip() != "":
                    non_empty_rows.append(cell.row)
                    non_empty_cols.append(cell.column)

        if non_empty_rows and non_empty_cols:
            min_r, max_r = min(non_empty_rows), max(non_empty_rows)
            min_c, max_c = min(non_empty_cols), max(non_empty_cols)

            # Only flag if there are at least 2 rows and 2 columns of data
            if (max_r - min_r + 1 >= 2) and (max_c - min_c + 1 >= 2):
                covered_by_table = any(
                    _is_coord_in_table(min_r, min_c, max_r, max_c, t) for t in tables
                )
                if not covered_by_table:
                    ref_str = f"{get_column_letter(min_c)}{min_r}:{get_column_letter(max_c)}{max_r}"
                    loc = f"{ws.title}!{ref_str}"
                    findings.append(
                        Finding(
                            rule_id="table-header-missing",
                            sc="1.3.1",
                            severity="serious",
                            location=loc,
                            description=(
                                f"Data range '{ref_str}' is not structured as an official Excel Table "
                                "with designated column headers"
                            ),
                            evidence=f"range='{ref_str}'",
                            fixable=True,
                            fix="Convert range to an official Excel Table with header row",
                        )
                    )
    return findings


def check_charts_alt(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.1.1: Verifies embedded charts have alternative text or titles."""
    findings = []
    for ws in wb.worksheets:
        charts = getattr(ws, "_charts", [])
        for idx, chart in enumerate(charts):
            title_text = ""
            if chart.title:
                if isinstance(chart.title, str):
                    title_text = chart.title.strip()
                elif hasattr(chart.title, "text") and chart.title.text:
                    title_text = str(chart.title.text).strip()

            if not title_text:
                findings.append(
                    Finding(
                        rule_id="chart-alt-missing",
                        sc="1.1.1",
                        severity="serious",
                        location=f"{ws.title}!chart[{idx}]",
                        description=f"Chart #{idx + 1} in sheet '{ws.title}' lacks a title or alternative description",
                        evidence="chart.title is empty",
                        fixable=False,
                        fix="Add a descriptive title or alternative text to the chart",
                    )
                )
    return findings


def check_images_alt(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.1.1: Verifies embedded images have alternative text."""
    findings = []
    generic_alts = {"picture", "image", "graphic", "photo"}
    for ws in wb.worksheets:
        images = getattr(ws, "_images", [])
        for idx, img in enumerate(images):
            descr = ""
            title = ""
            # Direct attributes if set programmatically
            if hasattr(img, "descr") and img.descr:
                descr = str(img.descr).strip()
            if hasattr(img, "title") and img.title:
                title = str(img.title).strip()
            if hasattr(img, "alt_text") and img.alt_text:
                descr = str(img.alt_text).strip()

            # Inspect openpyxl Drawing anchor pic hierarchy
            if not descr and not title and hasattr(img, "anchor"):
                anchor = img.anchor
                if hasattr(anchor, "pic") and hasattr(anchor.pic, "nvPicPr"):
                    cNvPr = getattr(anchor.pic.nvPicPr, "cNvPr", None)
                    if cNvPr:
                        descr = (getattr(cNvPr, "descr", "") or "").strip()
                        title = (getattr(cNvPr, "title", "") or "").strip()

            # Underlying XML element if available
            if not descr and not title and hasattr(img, "_element") and img._element is not None:
                cNvPr = img._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr")
                if cNvPr is not None:
                    descr = (cNvPr.get("descr") or "").strip()
                    title = (cNvPr.get("title") or "").strip()

            # Generic placeholders like "Picture" count as missing alt
            if descr.lower() in generic_alts:
                descr = ""

            if not descr and not title:
                anchor_loc = getattr(img, "anchor", None)
                loc_str = f"image[{idx}]"
                if isinstance(anchor_loc, str):
                    loc_str = f"image[{anchor_loc}]"
                findings.append(
                    Finding(
                        rule_id="image-alt-missing",
                        sc="1.1.1",
                        severity="serious",
                        location=f"{ws.title}!{loc_str}",
                        description=f"Embedded image #{idx + 1} in sheet '{ws.title}' lacks alternative text",
                        evidence="cNvPr descr/title is empty or generic placeholder",
                        fixable=False,
                        fix="Add an alternative text description to the image",
                    )
                )
    return findings


def check_color_contrast(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 1.4.3: Evaluates text vs background contrast in populated cells."""
    findings = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue

                # Background color determination
                bg_rgb = (1.0, 1.0, 1.0)  # default white
                if cell.fill and cell.fill.fill_type:
                    fg = cell.fill.fgColor
                    if fg and fg.rgb:
                        parsed = hex_to_rgb_tuple(str(fg.rgb))
                        if parsed:
                            bg_rgb = parsed

                # Font color determination
                fg_rgb = (0.0, 0.0, 0.0)  # default black
                font = cell.font
                if font and font.color and font.color.rgb:
                    parsed = hex_to_rgb_tuple(str(font.color.rgb))
                    if parsed:
                        fg_rgb = parsed

                ratio = wcag_contrast_ratio.rgb(fg_rgb, bg_rgb)

                # Threshold: 3.0:1 for large text (>= 18pt or >= 14pt bold), else 4.5:1
                size = getattr(font, "size", 11) or 11
                bold = bool(getattr(font, "bold", False))
                is_large = (size >= 18) or (bold and size >= 14)
                threshold = 3.0 if is_large else 4.5

                if ratio < threshold:
                    findings.append(
                        Finding(
                            rule_id="color-contrast",
                            sc="1.4.3",
                            severity="moderate",
                            location=f"{ws.title}!{cell.coordinate}",
                            description=(
                                f"Cell '{cell.coordinate}' text contrast ratio {ratio:.2f}:1 "
                                f"is below the required {threshold:.1f}:1 threshold"
                            ),
                            evidence=f"ratio={ratio:.2f}, size={size}, bold={bold}",
                            fixable=True,
                            fix="Adjust font color or cell fill to meet WCAG contrast ratio",
                        )
                    )
    return findings


def check_hyperlinks(wb: openpyxl.Workbook) -> List[Finding]:
    """WCAG 2.4.4: Flags ambiguous link text or raw URLs."""
    findings = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.hyperlink and cell.hyperlink.target:
                    target = str(cell.hyperlink.target).strip()
                    val = str(cell.value or "").strip()
                    val_lower = val.lower()

                    if val_lower in VAGUE_LINK_TEXTS or val == target:
                        findings.append(
                            Finding(
                                rule_id="link-text-vague",
                                sc="2.4.4",
                                severity="moderate",
                                location=f"{ws.title}!{cell.coordinate}",
                                description=f"Hyperlink in cell '{cell.coordinate}' has vague or raw URL text '{val}'",
                                evidence=f"text='{val}', target='{target}'",
                                fixable=False,
                                fix="Replace link text with a descriptive phrase explaining the target",
                            )
                        )
    return findings


ALL_RULES = [
    check_workbook_title,
    check_sheet_names,
    check_empty_sheets,
    check_merged_cells,
    check_table_headers,
    check_charts_alt,
    check_images_alt,
    check_color_contrast,
    check_hyperlinks,
]


def audit_rules(wb: openpyxl.Workbook) -> List[Finding]:
    """Runs all audit rules on an openpyxl Workbook."""
    all_findings = []
    for rule_fn in ALL_RULES:
        all_findings.extend(rule_fn(wb))
    return all_findings
