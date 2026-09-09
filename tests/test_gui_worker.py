"""Tests for GUI batch worker thread."""
from pathlib import Path
import openpyxl
from xslx_a11y.gui.models import BatchItem
from xslx_a11y.gui.worker import BatchWorker


def test_batch_worker_run(tmp_path: Path, qtbot):
    sample = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    wb.properties.title = "Sample Workbook"
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Metric"
    ws["B1"] = "Score"
    from openpyxl.worksheet.table import Table
    ws.add_table(Table(displayName="MetricTable", ref="A1:B1"))
    wb.save(sample)

    out_dir = tmp_path / "out"
    item = BatchItem(path=sample)
    worker = BatchWorker(
        items=[item],
        out_dir=out_dir,
        formats=["md", "json"],
        theme="light",
        auto_fix=False,
    )

    with qtbot.waitSignal(worker.all_completed, timeout=5000):
        worker.start()

    assert item.status == "Completed"
    assert item.findings_count == 0
    assert item.score == 100.0
    assert "md" in item.reports
    assert item.reports["md"].exists()
