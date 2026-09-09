"""Tests for criteria checklist CLI flags in xslx-a11y (--init-criteria and --criteria)."""
from pathlib import Path
import openpyxl
import pytest
from xslx_a11y.cli import main


def test_cli_init_criteria(tmp_path: Path):
    target = tmp_path / "criteria.md"
    with pytest.raises(SystemExit) as exc:
        main(["--init-criteria", str(target)])
    assert exc.value.code == 0
    assert target.exists()
    content = target.read_text(encoding="utf-8")
    assert "Criteria Checklist" in content
    assert "[x] 1.1.1" in content


def test_cli_criteria_what_if(tmp_path: Path):
    # Create a workbook with violations
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Overview"  # non-default sheet name
    ws.merge_cells("A1:B1")
    ws["A1"] = "Merged Header"
    src = tmp_path / "violations.xlsx"
    wb.save(src)

    criteria_file = tmp_path / "criteria.md"
    criteria_file.write_text(
        "# Custom Criteria\n\n"
        "[ ] 1.1.1 Non-text Content\n"
        "[ ] 1.3.1 Info and Relationships\n"
        "[ ] 2.4.2 Page Titled\n",
        encoding="utf-8",
    )
    out_dir = tmp_path / "reports"
    with pytest.raises(SystemExit) as exc:
        main(
            [
                str(src),
                "--format",
                "md",
                "--output-dir",
                str(out_dir),
                "--criteria",
                str(criteria_file),
            ]
        )
    # With violations partitioned to excluded, should exit 0 instead of failing with 1
    assert exc.value.code == 0
    report_file = out_dir / "violations-a11y-report.md"
    assert report_file.exists()
    content = report_file.read_text(encoding="utf-8")
    assert "What-If" in content or "EXCLUDED" in content
