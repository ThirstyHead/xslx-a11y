"""Tests for workbook metadata rules (e.g. title-missing)."""
import openpyxl
from xlsx_a11y.rules import check_workbook_title


def test_title_missing_when_empty():
    wb = openpyxl.Workbook()
    wb.properties.title = None
    findings = check_workbook_title(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "title-missing"
    assert findings[0].severity == "critical"
    assert findings[0].fixable is True


def test_title_missing_when_whitespace():
    wb = openpyxl.Workbook()
    wb.properties.title = "   "
    findings = check_workbook_title(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "title-missing"


def test_title_present():
    wb = openpyxl.Workbook()
    wb.properties.title = "Quarterly Financial Accessibility Report"
    findings = check_workbook_title(wb)
    assert len(findings) == 0
