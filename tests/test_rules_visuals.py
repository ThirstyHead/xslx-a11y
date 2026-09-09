"""Tests for visual element rules (images and charts alt text)."""
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
from xslx_a11y.rules import check_images_alt, check_charts_alt


def test_chart_alt_missing_when_no_title(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Charts"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    ws.append(["B", 20])

    chart = BarChart()
    chart.title = None
    data = Reference(ws, min_col=2, min_row=1, max_row=3)
    chart.add_data(data, titles_from_data=True)
    ws.add_chart(chart, "D2")

    findings = check_charts_alt(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "chart-alt-missing"
    assert findings[0].severity == "serious"


def test_chart_alt_present(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Charts"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])

    chart = BarChart()
    chart.title = "Revenue by Category"
    ws.add_chart(chart, "D2")

    findings = check_charts_alt(wb)
    assert len(findings) == 0


def test_image_alt_missing(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visuals"

    # Create dummy 10x10 image
    img_path = tmp_path / "test.png"
    img = PILImage.new("RGB", (10, 10), color="blue")
    img.save(img_path)

    xl_img = Image(str(img_path))
    ws.add_image(xl_img, "B2")

    findings = check_images_alt(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "image-alt-missing"
    assert findings[0].severity == "serious"


def test_image_alt_present(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visuals"

    img_path = tmp_path / "test.png"
    img = PILImage.new("RGB", (10, 10), color="blue")
    img.save(img_path)

    xl_img = Image(str(img_path))
    xl_img.descr = "Company Logo in Blue"

    ws.add_image(xl_img, "B2")

    findings = check_images_alt(wb)
    assert len(findings) == 0
