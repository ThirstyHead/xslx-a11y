"""Tests for interactive triage workflow."""
from pathlib import Path
import openpyxl
from xslx_a11y.triage import run_interactive_triage


def test_interactive_triage_terminal(tmp_path: Path):
    src = tmp_path / "triage_input.xlsx"
    wb = openpyxl.Workbook()
    wb.properties.title = None
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Data"
    wb.save(src)

    out = tmp_path / "triaged_output.xlsx"

    # User inputs for: title prompt -> "My Financials"
    responses = iter(["My Financials", "s"])
    messages = []

    def mock_input(prompt=""):
        return next(responses, "s")

    def mock_print(*args, **kwargs):
        messages.append(" ".join(str(a) for a in args))

    items = run_interactive_triage(
        in_path=src,
        out_path=out,
        input_func=mock_input,
        print_func=mock_print,
    )

    assert items >= 1
    assert out.exists()
    triaged_wb = openpyxl.load_workbook(out)
    assert triaged_wb.properties.title == "My Financials"
