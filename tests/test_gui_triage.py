"""Tests for interactive GUI triage dialog."""
from pathlib import Path
import openpyxl
from xslx_a11y.findings import Finding
from xslx_a11y.gui.triage_dialog import TriageDialog


def test_triage_dialog_flow(tmp_path: Path, qtbot):
    sample = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    wb.save(sample)

    findings = [
        Finding(
            rule_id="title-missing",
            sc="2.4.2",
            severity="critical",
            location="core.xml",
            description="Document title missing",
            evidence="",
            fixable=True,
        ),
        Finding(
            rule_id="sheet-name-default",
            sc="2.4.6",
            severity="moderate",
            location="Sheet1",
            description="Default sheet name",
            evidence="",
            fixable=False,
        ),
    ]

    dialog = TriageDialog(sample, findings)
    qtbot.addWidget(dialog)

    # 1. Provide title
    assert dialog.lbl_progress.text() == "Barrier 1 of 2"
    dialog.txt_input.setText("Audited Financials")
    dialog.apply_current()

    # 2. Rename sheet
    assert dialog.lbl_progress.text() == "Barrier 2 of 2"
    dialog.txt_input.setText("Q1_Overview")
    dialog.apply_current()

    # Verify triaged file was generated
    triaged_p = tmp_path / "test-triaged.xlsx"
    assert triaged_p.exists()

    reloaded = openpyxl.load_workbook(triaged_p)
    assert reloaded.properties.title == "Audited Financials"
    assert "Q1_Overview" in reloaded.sheetnames
