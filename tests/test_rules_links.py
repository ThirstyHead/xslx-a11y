"""Tests for vague hyperlink text rules."""
import openpyxl
from xslx_a11y.rules import check_hyperlinks


def test_vague_link_text():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"

    cell = ws["A1"]
    cell.value = "Click here"
    cell.hyperlink = "https://example.com/report"

    findings = check_hyperlinks(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "link-text-vague"
    assert findings[0].severity == "moderate"


def test_bare_url_link_text():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"

    cell = ws["A1"]
    cell.value = "https://example.com/report"
    cell.hyperlink = "https://example.com/report"

    findings = check_hyperlinks(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "link-text-vague"


def test_descriptive_link_text():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"

    cell = ws["A1"]
    cell.value = "Annual Accessibility Report 2026"
    cell.hyperlink = "https://example.com/report"

    findings = check_hyperlinks(wb)
    assert len(findings) == 0
