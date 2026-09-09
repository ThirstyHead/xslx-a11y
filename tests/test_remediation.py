"""Tests for deterministic remediation engine."""
from pathlib import Path
import openpyxl
import pytest
from xlsx_a11y.immutability import calculate_sha256
from xlsx_a11y.remediate import remediate_file, remediate_workbook


def test_remediate_workbook_fixes_title_merged_and_tables():
    wb = openpyxl.Workbook()
    wb.properties.title = None
    ws = wb.active
    ws.title = "Summary"

    # Merged cell
    ws.merge_cells("A1:C1")
    ws["A1"] = "Consolidated Sales"

    # Data block
    ws["A2"] = "Region"
    ws["B2"] = "Q1"
    ws["C2"] = "Q2"
    ws["A3"] = "North"
    ws["B3"] = 100
    ws["C3"] = 200

    remediated_wb, fixes = remediate_workbook(wb, fallback_title="Quarterly Sales")

    # Title should be set
    assert remediated_wb.properties.title == "Quarterly Sales"

    # Cells should be unmerged and value replicated
    assert len(ws.merged_cells.ranges) == 0
    assert ws["A1"].value == "Consolidated Sales"
    assert ws["B1"].value == "Consolidated Sales"
    assert ws["C1"].value == "Consolidated Sales"

    # Official table should be created
    assert len(ws.tables) >= 1
    assert any("title" in f.lower() for f in fixes)
    assert any("unmerge" in f.lower() for f in fixes)


def test_remediate_file_end_to_end_immutability(tmp_path: Path):
    src = tmp_path / "Quarterly-Report.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.merge_cells("B2:D2")
    ws["B2"] = "Header"
    ws["B3"] = 10
    ws["C3"] = 20
    ws["D3"] = 30
    wb.save(src)

    sha_before = calculate_sha256(src)
    out = tmp_path / "Quarterly-Report-remediated.xlsx"

    res = remediate_file(src, out)

    sha_after = calculate_sha256(src)
    assert sha_before == sha_after, "Source file was modified!"
    assert res["original_file_immutable"] is True
    assert out.exists()
    assert res["remediated_sha256"] == calculate_sha256(out)
    assert len(res["remediations_applied"]) > 0


def test_remediate_file_rejects_same_path(tmp_path: Path):
    src = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    wb.save(src)

    with pytest.raises(ValueError, match="strictly guarantees"):
        remediate_file(src, src)
