"""Tests for audit orchestrator."""
from pathlib import Path
import openpyxl
from xslx_a11y.audit import audit_workbook, audit_file


def test_audit_workbook_clean():
    wb = openpyxl.Workbook()
    wb.properties.title = "Clean Workbook"
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Key"
    ws["B1"] = "Value"
    ws["A2"] = "Metric"
    ws["B2"] = 100

    from openpyxl.worksheet.table import Table
    tab = Table(displayName="OverviewTable", ref="A1:B2")
    ws.add_table(tab)

    result = audit_workbook(wb, file_path="dummy.xlsx")
    assert result["summary"]["pass"] is True
    assert result["summary"]["blocking"] == 0
    assert result["summary"]["total"] == 0
    assert len(result["findings"]) == 0


def test_audit_file(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"  # default sheet name
    ws.merge_cells("B2:C2")
    ws["B2"] = "Merged"

    p = tmp_path / "sample.xlsx"
    wb.save(p)

    result = audit_file(p)
    assert result["file"] == str(p.resolve())
    assert len(result["sha256"]) == 64
    assert "audited_at" in result
    assert result["summary"]["pass"] is False
    assert result["summary"]["blocking"] >= 1  # merged-cell or title-missing

    rule_ids = {f.rule_id for f in result["findings"]}
    assert "title-missing" in rule_ids
    assert "sheet-name-default" in rule_ids
    assert "merged-cell" in rule_ids
