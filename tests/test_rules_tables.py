"""Tests for table headers and structured data rules."""
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from xslx_a11y.rules import check_table_headers


def test_table_header_missing_when_raw_data_block():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    # Create 3x3 data block without Table object
    ws["A1"] = "Month"
    ws["B1"] = "Revenue"
    ws["C1"] = "Cost"
    ws["A2"] = "Jan"
    ws["B2"] = 100
    ws["C2"] = 50
    ws["A3"] = "Feb"
    ws["B3"] = 120
    ws["C3"] = 60

    findings = check_table_headers(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "table-header-missing"
    assert findings[0].severity == "serious"
    assert "A1:C3" in findings[0].location
    assert findings[0].fixable is True


def test_table_with_headers_passes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Month"
    ws["B1"] = "Revenue"
    ws["A2"] = "Jan"
    ws["B2"] = 100

    tab = Table(displayName="SalesTable", ref="A1:B2")
    tab.headerRowCount = 1
    ws.add_table(tab)

    findings = check_table_headers(wb)
    assert len(findings) == 0


def test_table_with_header_row_count_zero_fails():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Jan"
    ws["B1"] = 100

    tab = Table(displayName="SalesTableNoHeader", ref="A1:B1")
    tab.headerRowCount = 0
    ws.add_table(tab)

    findings = check_table_headers(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "table-header-missing"
