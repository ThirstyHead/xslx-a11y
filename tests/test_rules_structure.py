"""Tests for workbook structure rules (default names, empty sheets, merged cells)."""
import openpyxl
from xlsx_a11y.rules import (
    check_sheet_names,
    check_empty_sheets,
    check_merged_cells,
)


def test_sheet_name_default():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    findings = check_sheet_names(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "sheet-name-default"
    assert findings[0].location == "Sheet1"
    assert findings[0].severity == "moderate"

    # Meaningful name
    ws.title = "Revenue_2026"
    assert len(check_sheet_names(wb)) == 0


def test_sheet_empty():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BlankSheet"
    findings = check_empty_sheets(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "sheet-empty"
    assert findings[0].location == "BlankSheet"

    # Populate cell
    ws["A1"] = "Data point"
    assert len(check_empty_sheets(wb)) == 0


def test_merged_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.merge_cells("B2:D2")
    ws["B2"] = "Header Group"

    findings = check_merged_cells(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "merged-cell"
    assert findings[0].severity == "serious"
    assert "B2:D2" in findings[0].location
    assert findings[0].fixable is True
