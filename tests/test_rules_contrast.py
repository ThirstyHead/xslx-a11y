"""Tests for color contrast rules in worksheets."""
import openpyxl
from openpyxl.styles import Font, PatternFill
from xlsx_a11y.rules import check_color_contrast, check_red_formatting


def test_contrast_failure_light_gray_on_white():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contrast"

    # Light gray text on default white background
    cell = ws["A1"]
    cell.value = "Low contrast text"
    cell.font = Font(color="FFAAAAAA", size=11)
    cell.fill = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")

    findings = check_color_contrast(wb)
    assert len(findings) == 1
    assert findings[0].rule_id == "color-contrast"
    assert findings[0].severity == "moderate"
    assert "A1" in findings[0].location


def test_contrast_pass_black_on_white():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contrast"

    cell = ws["A1"]
    cell.value = "High contrast text"
    cell.font = Font(color="FF000000", size=11)
    cell.fill = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")

    findings = check_color_contrast(wb)
    assert len(findings) == 0


def test_contrast_large_text_threshold():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contrast"

    # Gray color ~3.5:1 contrast
    # #767676 on #FFFFFF is ~4.54:1 (passes 4.5)
    # #888888 on #FFFFFF is ~3.5:1 (fails normal 4.5:1, passes large 3.0:1)
    cell = ws["A1"]
    cell.value = "Large Heading Text"
    cell.font = Font(color="FF888888", size=18, bold=False)
    cell.fill = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")

    findings = check_color_contrast(wb)
    assert len(findings) == 0


def test_red_formatting_detection():
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "RedCheck"

    # Red font
    ws["A1"] = "Negative variance"
    ws["A1"].font = Font(color="FFFF0000")

    # Red fill
    ws["B1"] = "Alert"
    ws["B1"].fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

    # Red number format
    ws["C1"] = 1234
    ws["C1"].number_format = "$#,##0;[Red]($#,##0)"

    # Clean cell
    ws["D1"] = "Normal text"

    findings = check_red_formatting(wb)
    assert len(findings) == 3
    coords = [f.location for f in findings]
    assert "RedCheck!A1" in coords
    assert "RedCheck!B1" in coords
    assert "RedCheck!C1" in coords
    for f in findings:
        assert f.rule_id == "color-use-red"
        assert f.sc == "1.4.1"
