"""End-to-end tests for xlsx-a11y CLI."""
from pathlib import Path
import openpyxl
import pytest
from xlsx_a11y.cli import main


def test_cli_missing_file():
    with pytest.raises(SystemExit) as exc:
        main(["nonexistent.xlsx"])
    assert exc.value.code == 2


def test_cli_audit_clean_file(tmp_path: Path):
    wb = openpyxl.Workbook()
    wb.properties.title = "Clean Summary"
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Key"
    ws["B1"] = "Value"
    ws["A2"] = "Total"
    ws["B2"] = 100

    from openpyxl.worksheet.table import Table
    tab = Table(displayName="OverviewTable", ref="A1:B2")
    ws.add_table(tab)

    src = tmp_path / "clean.xlsx"
    wb.save(src)

    out_dir = tmp_path / "reports"
    with pytest.raises(SystemExit) as exc:
        main([str(src), "--output-dir", str(out_dir), "--format", "md,html,pdf,json"])
    assert exc.value.code == 0

    assert (out_dir / "clean-a11y-report.md").exists()
    assert (out_dir / "clean-a11y-report.html").exists()
    assert (out_dir / "clean-a11y-report.pdf").exists()
    assert (out_dir / "clean-audit.json").exists()


def test_cli_fix_flag(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:B1")
    ws["A1"] = "Merged"

    src = tmp_path / "unfixed.xlsx"
    wb.save(src)

    out_dir = tmp_path / "reports"
    fixed_xlsx = tmp_path / "fixed.xlsx"

    with pytest.raises(SystemExit) as exc:
        main([str(src), "--fix", "--out-xlsx", str(fixed_xlsx), "--output-dir", str(out_dir)])
    assert fixed_xlsx.exists()
